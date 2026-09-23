"""Build CXO briefing DOCX for FinnAI SLM — plain, human, no ops fluff."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent
IMG = OUT_DIR / "_figures"
IMG.mkdir(exist_ok=True)


def _font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    return h


def para(doc, text, size=11, bold=False, after=8):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _font(r, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(0)
    return p


def note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _font(r, size=10, italic=True, color=(0x44, 0x44, 0x44))
    p.paragraph_format.space_after = Pt(10)


def add_img(doc, path: Path, width=6.0, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        _font(r, size=9, italic=True, color=(0x66, 0x66, 0x66))


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        _font(r, size=10, bold=True, color=(0xFF, 0xFF, 0xFF))
        sh = OxmlElement("w:shd")
        sh.set(qn("w:fill"), "1A1A1A")
        cell._tc.get_or_add_tcPr().append(sh)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            _font(r, size=10)
    doc.add_paragraph()


def fig_pipeline():
    fig, ax = plt.subplots(figsize=(10.5, 3.6))
    ax.set_xlim(0, 10.5)
    ax.set_ylim(0, 3.6)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    steps = [
        (0.25, "Build\ntraining data\n(synthetic)"),
        (2.3, "Split\ntrain / val / test"),
        (4.35, "Fine-tune\nsmall model\n(QLoRA)"),
        (6.4, "Score against\nfixed gates"),
        (8.45, "Ship on phone\n+ publish"),
    ]
    colors = ["#E8F0EE", "#E8F0EE", "#D4E5E1", "#E8F0EE", "#D4E5E1"]
    for (x, text), color in zip(steps, colors):
        ax.add_patch(
            FancyBboxPatch(
                (x, 0.7), 1.85, 2.2, boxstyle="round,pad=0.03,rounding_size=0.12",
                facecolor=color, edgecolor="#333", lw=1.0,
            )
        )
        ax.text(x + 0.925, 1.8, text, ha="center", va="center", fontsize=9)
    for i in range(4):
        x0 = steps[i][0] + 1.85
        x1 = steps[i + 1][0]
        ax.annotate("", xy=(x1, 1.8), xytext=(x0, 1.8),
                    arrowprops=dict(arrowstyle="->", color="#222", lw=1.4))
    ax.set_title("What we actually did", fontsize=12, pad=6)
    fig.tight_layout()
    path = IMG / "pipeline.png"
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close()
    return path


def fig_qlora():
    fig, ax = plt.subplots(figsize=(9.5, 3.8))
    ax.set_xlim(0, 9.5)
    ax.set_ylim(0, 3.8)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.add_patch(FancyBboxPatch((0.3, 0.6), 4.0, 2.6, boxstyle="round,pad=0.04,rounding_size=0.12",
                                facecolor="#EEF3F8", edgecolor="#333"))
    ax.text(2.3, 2.7, "Base model (frozen)", ha="center", fontsize=11, fontweight="bold")
    ax.text(2.3, 1.7, "Qwen3 · 1.7B params\nAlready knows language.\nWe do not retrain all of it.",
            ha="center", va="center", fontsize=9)
    ax.add_patch(FancyBboxPatch((5.2, 0.6), 4.0, 2.6, boxstyle="round,pad=0.04,rounding_size=0.12",
                                facecolor="#F7F1E8", edgecolor="#333"))
    ax.text(7.2, 2.7, "Small adapter (trained)", ha="center", fontsize=11, fontweight="bold")
    ax.text(7.2, 1.7, "~67 MB of new weights\nLearns bank SMS → JSON\nand grounded money answers.",
            ha="center", va="center", fontsize=9)
    ax.annotate("", xy=(5.2, 1.9), xytext=(4.3, 1.9),
                arrowprops=dict(arrowstyle="<->", color="#222", lw=1.6))
    ax.set_title("Fine-tuning without rebuilding the whole model", fontsize=12, pad=6)
    fig.tight_layout()
    path = IMG / "qlora.png"
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close()
    return path


def fig_architecture():
    fig, ax = plt.subplots(figsize=(9.5, 4.4))
    ax.set_xlim(0, 9.5)
    ax.set_ylim(0, 4.4)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.add_patch(FancyBboxPatch((0.25, 0.25), 9.0, 3.9, boxstyle="round,pad=0.02,rounding_size=0.1",
                                facecolor="#FAFAFA", edgecolor="#222", lw=1.2))
    ax.text(4.75, 3.85, "On the phone", ha="center", fontsize=11, fontweight="bold")
    ax.add_patch(FancyBboxPatch((0.5, 2.6), 2.4, 0.9, boxstyle="round,pad=0.02,rounding_size=0.08",
                                facecolor="#E8F0EE", edgecolor="#333"))
    ax.text(1.7, 3.05, "Bank SMS", ha="center", fontsize=10)
    ax.add_patch(FancyBboxPatch((3.3, 2.6), 3.0, 0.9, boxstyle="round,pad=0.02,rounding_size=0.08",
                                facecolor="#E8F0EE", edgecolor="#333"))
    ax.text(4.8, 3.05, "Known banks →\nregex parsers", ha="center", fontsize=9)
    ax.add_patch(FancyBboxPatch((6.7, 2.6), 2.3, 0.9, boxstyle="round,pad=0.02,rounding_size=0.08",
                                facecolor="#E8F0EE", edgecolor="#333"))
    ax.text(7.85, 3.05, "Done", ha="center", fontsize=10)
    ax.annotate("", xy=(3.3, 3.05), xytext=(2.9, 3.05), arrowprops=dict(arrowstyle="->", color="#222"))
    ax.annotate("", xy=(6.7, 3.05), xytext=(6.3, 3.05), arrowprops=dict(arrowstyle="->", color="#222"))
    ax.add_patch(FancyBboxPatch((2.2, 0.7), 5.1, 1.4, boxstyle="round,pad=0.03,rounding_size=0.1",
                                facecolor="#F7F1E8", edgecolor="#222", lw=1.2))
    ax.text(4.75, 1.7, "FinnAI SLM", ha="center", fontsize=11, fontweight="bold")
    ax.text(4.75, 1.15, "Unknown banks, Indic SMS, Ask/Learn coaching\nRuns locally. No SMS sent to a cloud LLM.",
            ha="center", fontsize=9)
    ax.annotate("", xy=(4.75, 2.1), xytext=(4.8, 2.6),
                arrowprops=dict(arrowstyle="->", color="#222", lw=1.2))
    ax.text(5.6, 2.35, "else", fontsize=8)
    ax.set_title("How it sits in the product", fontsize=12, pad=6)
    fig.tight_layout()
    path = IMG / "architecture.png"
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close()
    return path


def fig_results():
    fig, ax = plt.subplots(figsize=(8.5, 4.0))
    fig.patch.set_facecolor("white")
    labels = ["Base\nQwen3", "Old app\nbaseline", "FinnAI v1", "FinnAI v2"]
    vals = [28.2, 42.4, 92.8, 98.0]
    colors = ["#B0B0B0", "#909090", "#5F8A82", "#1A1A1A"]
    bars = ax.bar(labels, vals, color=colors, width=0.62)
    ax.set_ylabel("SMS fields fully correct %")
    ax.set_ylim(0, 110)
    ax.set_title("Held-out test — did fine-tuning work?", fontsize=12)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + 2, f"{v:.0f}%", ha="center", fontsize=10)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    path = IMG / "results.png"
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close()
    return path


def fig_data():
    fig, ax = plt.subplots(figsize=(7.5, 3.8))
    fig.patch.set_facecolor("white")
    sizes = [60, 25, 15]
    labels = ["SMS → JSON (60%)", "Finance coach (25%)", "General money Q&A (15%)"]
    colors = ["#5F8A82", "#8AA8C0", "#C4A574"]
    ax.pie(sizes, labels=labels, colors=colors, startangle=90,
           wedgeprops=dict(width=0.5, edgecolor="white"))
    ax.set_title("What we trained on (16k examples)", fontsize=12)
    fig.tight_layout()
    path = IMG / "data_mix.png"
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close()
    return path


def build(paths):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("FinnAI SLM")
    _font(r, size=26, bold=True)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("How we taught a phone-sized model to read bank SMS")
    _font(r, size=13, color=(0x33, 0x33, 0x33))

    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = m.add_run(
        "FinnDot · 2026\n"
        "Model — huggingface.co/finndot/finnai-slm-v2\n"
        "Data — huggingface.co/datasets/finndot/finnai-slm-data\n"
        "Code — github.com/arunbhatg/finnai-slm"
    )
    _font(r, size=9, color=(0x66, 0x66, 0x66))
    doc.add_paragraph()

    # ---
    heading(doc, "In short")
    para(
        doc,
        "FinnDot reads bank SMS and logs expenses. For banks we already know, we use hand-written parsers. "
        "For everything else — new formats, Hindi/Hinglish texts, and “where did my money go?” questions — "
        "we needed a small language model that runs on the phone.",
    )
    para(
        doc,
        "We started from an open model (Qwen3, 1.7B parameters) and fine-tuned it on synthetic data only. "
        "No real user SMS. On a held-out test, full-field SMS accuracy went from about 28% (untuned) to about 98% (FinnAI v2).",
    )
    add_img(doc, paths["results"], caption="Figure 1. Same test set, four models.")

    # ---
    heading(doc, "The problem we were solving")
    para(
        doc,
        "A bank text looks like this:",
    )
    note(doc, "HDFC Bank: Rs.499.00 debited from A/c XX1234 to SWIGGY via UPI. Avl Bal Rs.15000.50")
    para(doc, "We need this:")
    note(
        doc,
        '{"amount": 499.0, "merchant": "SWIGGY", "type": "EXPENSE", "account": "1234", "balance": 15000.5}',
    )
    para(
        doc,
        "Rules work great when the format is known. They break when the bank changes wording, "
        "the message is in another language, or the text is an OTP and should be ignored. "
        "Coaching is not a regex problem at all.",
    )
    para(
        doc,
        "Cloud chat APIs could do the job, but bank SMS has balances and account fragments. "
        "We did not want that traffic leaving the device for this path. Local model, download once, run offline.",
    )

    # ---
    heading(doc, "What “fine-tuning” means here")
    para(
        doc,
        "If you only use ChatGPT as a product: fine-tuning is not “make a new ChatGPT.” "
        "It is taking a general model that already speaks language, and teaching it one job "
        "with many graded examples.",
    )
    para(
        doc,
        "We used QLoRA — keep the big model frozen, train a thin adapter (~67 MB). "
        "Cheaper than full retraining, less chance of wiping general ability, fits on a single modern GPU.",
    )
    add_img(doc, paths["qlora"], caption="Figure 2. Adapter on top of a frozen base.")

    heading(doc, "Settings that mattered", 2)
    table(
        doc,
        ["Choice", "What we used", "Why"],
        [
            ["Base", "Qwen3-1.7B", "Open license, small enough for phones"],
            ["Method", "QLoRA SFT", "Specialise without full retrain"],
            ["LoRA rank", "16", "Enough for JSON style; larger did not help"],
            ["Epochs", "3", "16k examples; more started to overfit patterns"],
            ["Pick checkpoint by", "Val SMS exact-match", "Loss alone lied to us before"],
            ["Thinking / chain-of-thought", "Off", "Phone needs a fast JSON answer"],
        ],
    )

    # ---
    heading(doc, "Where it sits in the app")
    add_img(doc, paths["architecture"], caption="Figure 3. Parsers first. Model for the rest.")
    para(
        doc,
        "We did not throw away the 35+ bank parsers. The model is the long-tail and the coach.",
    )

    # ---
    heading(doc, "How we did it")
    add_img(doc, paths["pipeline"], width=6.2, caption="Figure 4. End to end.")

    heading(doc, "Data", 2)
    para(
        doc,
        "All training SMS are synthetic: templates with fake amounts and merchants. "
        "We know the correct JSON because we generated the text. "
        "Coach answers are tied to fake ledgers so every rupee in the reply must appear in the prompt.",
    )
    add_img(doc, paths["data"], width=5.2, caption="Figure 5. Mix of tasks in training.")
    para(
        doc,
        "Splits are by template family, not by line. If you put paraphrases of the same template in both train and test, "
        "you fool yourself. We locked that before scoring.",
    )
    note(
        doc,
        "Indic paraphrases (Hindi, Hinglish, etc.) were generated as surface text only and kept only when "
        "amount, merchant, and account digits still matched our labels.",
    )

    heading(doc, "How we decided “good enough to ship”", 2)
    para(
        doc,
        "The pass rules were written before we looked at the final test numbers. "
        "Main score: every key field correct (amount, type, merchant, last-4, balance), or empty {} for non-transactions.",
    )
    table(
        doc,
        ["Check", "Bar", "v2"],
        [
            ["Beat untuned Qwen3 on SMS", "+3 points, clear gap", "Pass"],
            ["Do not get worse on amounts", "Within 1 point of base", "Pass"],
            ["Coach does not invent numbers", "Within 5 points of base", "Pass"],
            ["Valid JSON", "≥ 95%", "100%"],
            ["Beat old on-device baseline", "R-EM ≥ old model", "Pass"],
            ["Phone speed / memory", "Comparable to old model", "Still to measure on device"],
        ],
    )

    # ---
    heading(doc, "Numbers")
    table(
        doc,
        ["", "FinnAI v2", "v1", "Untuned Qwen3"],
        [
            ["SMS fully correct", "98.0%", "92.8%", "28.2%"],
            ["Amount correct", "99.5%", "94.6%", "82.8%"],
            ["Merchant correct", "98.2%", "93.1%", "61.3%"],
            ["OTP/promo wrongly parsed", "0.5%", "28.6%", "100%"],
            ["Coach grounded in ledger", "68%", "32%", "26%"],
        ],
    )
    para(
        doc,
        "The false-parse drop matters in product terms: fewer OTPs showing up as spends. "
        "Coach groundedness doubled from v1 to v2 — still not perfect, but usable.",
    )

    # ---
    heading(doc, "Cost")
    para(
        doc,
        "One training run is a few hours on a single GPU — low double-digit dollars for the whole v1→v2 loop "
        "including eval and Indic data generation. No annotation vendor. No user inbox dump.",
    )
    para(
        doc,
        "Ongoing: downloading the model once is cheaper than paying per SMS to a cloud API forever.",
    )

    # ---
    heading(doc, "What we published")
    para(
        doc,
        "Weights, dataset, training/eval code, and this write-up are public (Apache 2.0). "
        "Links are on the cover. Others can fine-tune further or copy the approach for other private on-device jobs.",
    )

    # ---
    heading(doc, "Honest limits")
    para(
        doc,
        "Built for Indian bank SMS. Regex still wins on banks we fully support. "
        "Merchant names are the fiddliest field. Coach answers can still paraphrase numbers sometimes. "
        "We still need a real-device speed/memory check before calling the phone build fully done.",
    )

    # ---
    heading(doc, "Next")
    para(
        doc,
        "Finish on-device timing, point the app at the new file when that passes, "
        "keep growing templates for new banks and languages — still without training on user SMS.",
    )

    heading(doc, "Glossary")
    table(
        doc,
        ["Term", "Meaning"],
        [
            ["LLM / SLM", "Large / small language model"],
            ["Fine-tune", "Extra training for one job"],
            ["QLoRA", "Train a small adapter; keep the base frozen"],
            ["SFT", "Learn from input → correct output pairs"],
            ["R-EM", "All key SMS fields match gold"],
            ["Grounded", "Numbers in the answer exist in the ledger prompt"],
        ],
    )

    out = OUT_DIR / "FinnAI_SLM_CXO_FineTuning_Brief.docx"
    doc.save(out)
    return out


def main():
    paths = {
        "pipeline": fig_pipeline(),
        "qlora": fig_qlora(),
        "architecture": fig_architecture(),
        "results": fig_results(),
        "data": fig_data(),
    }
    print(build(paths))


if __name__ == "__main__":
    main()
